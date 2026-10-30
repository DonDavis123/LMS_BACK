"""Read an uploaded CSV or .xlsx file into plain rows.

Pure Python plus ``openpyxl``: no Django models and no database access.

``parse_upload`` returns a ``ParsedFile`` whose rows are
``{"row_number": int, "values": {header: text}}``. ``row_number`` is the row
number the user sees in Excel (the first data row under a header on row 1 is
row 2). Problems with the file as a whole raise ``ImportFileError`` carrying a
code, a plain-English message and a how-to-fix hint.
"""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field

from .lead_import_cleaning import to_text as _to_text

_GUARD_CHARS = ("=", "+", "-", "@")


def to_text(raw: object) -> str:
    """A cell as text. A leading single quote in front of ``=``, ``+``, ``-`` or
    ``@`` is the guard our own CSV downloads add (and Excel's text prefix); it
    is not part of the value, so a downloaded rows-to-fix file uploads back
    unchanged."""
    text = _to_text(raw)
    if len(text) > 1 and text[0] == "'" and text[1] in _GUARD_CHARS:
        return text[1:]
    return text

DEFAULT_MAX_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_ROWS = 10_000
MAX_HEADER_ROW = 50
# Excel files often carry formatting down to row 1,048,576; stop reading
# after this many blank rows in a row instead of walking all of them.
MAX_CONSECUTIVE_BLANK_ROWS = 2_000

ALLOWED_EXTENSIONS = (".csv", ".xlsx")


class ImportFileError(Exception):
    """The uploaded file can't be imported as a whole."""

    def __init__(self, code: str, message: str, how_to_fix: str = ""):
        super().__init__(message)
        self.code = code
        self.message = message
        self.how_to_fix = how_to_fix

    def to_dict(self) -> dict:
        return {"code": self.code, "detail": self.message, "how_to_fix": self.how_to_fix}


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[dict]
    total_rows: int
    sheet_names: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    header_row: int = 1
    sheet_name: str | None = None


def _setting(name: str, default: int) -> int:
    """A Django setting if Django is configured, else the default."""
    try:
        from django.conf import settings

        return int(getattr(settings, name, default))
    except Exception:  # noqa: BLE001 - not configured (unit tests, scripts)
        return default


def _megabytes(size: int) -> str:
    return f"{size / (1024 * 1024):.0f} MB" if size >= 1024 * 1024 else f"{size // 1024} KB"


# --------------------------------------------------------------------------
# Public entry point
# --------------------------------------------------------------------------


def parse_upload(file_obj, filename: str, header_row: int = 1, sheet_name: str | None = None) -> ParsedFile:
    max_bytes = _setting("LEAD_IMPORT_MAX_BYTES", DEFAULT_MAX_BYTES)
    max_rows = _setting("LEAD_IMPORT_MAX_ROWS", DEFAULT_MAX_ROWS)

    extension = os.path.splitext(filename or "")[1].lower()
    if extension == ".xls":
        raise ImportFileError(
            "unsupported_xls",
            "Old Excel files (.xls) are not supported.",
            "Open the file in Excel, choose Save As, pick .xlsx or .csv, and upload that file.",
        )
    if extension not in ALLOWED_EXTENSIONS:
        shown = extension or "no extension"
        raise ImportFileError(
            "unsupported_file_type",
            f"Files of type {shown} can't be imported. Only .csv and .xlsx files are accepted.",
            "Save your data as a .csv or .xlsx file and upload that file.",
        )

    if not isinstance(header_row, int) or isinstance(header_row, bool) or not 1 <= header_row <= MAX_HEADER_ROW:
        raise ImportFileError(
            "invalid_header_row",
            f"The header row must be a number between 1 and {MAX_HEADER_ROW}.",
            "Enter the row number that holds your column names (usually 1).",
        )

    declared = getattr(file_obj, "size", None)
    if isinstance(declared, int) and declared > max_bytes:
        raise _too_large(declared, max_bytes)
    data = _read_bytes(file_obj, max_bytes)
    if not data.strip():
        raise ImportFileError(
            "empty_file",
            "The file is empty.",
            "Check that you chose the right file and that it contains a header row and data.",
        )

    if extension == ".csv":
        raw_rows, sheet_names, chosen_sheet = _read_csv(data), [], None
    else:
        raw_rows, sheet_names, chosen_sheet = _read_xlsx(data, header_row, sheet_name, max_rows)

    return _build(raw_rows, header_row, max_rows, sheet_names, chosen_sheet)


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------


def _too_large(size: int, max_bytes: int) -> ImportFileError:
    return ImportFileError(
        "file_too_large",
        f"The file is {_megabytes(size)}; the limit is {_megabytes(max_bytes)}.",
        "Split the file into smaller files, or remove unused columns and sheets, and upload again.",
    )


def _read_bytes(file_obj, max_bytes: int) -> bytes:
    try:
        if isinstance(file_obj, (bytes, bytearray)):
            data = bytes(file_obj)
        else:
            if hasattr(file_obj, "seek"):
                file_obj.seek(0)
            data = file_obj.read(max_bytes + 1)
            if isinstance(data, str):
                data = data.encode("utf-8")
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError(
            "unreadable_file",
            "The file could not be read.",
            "Upload the file again. If it keeps failing, save a fresh copy and try that one.",
        ) from exc
    if len(data) > max_bytes:
        raise _too_large(len(data), max_bytes)
    return data


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")  # never fails


def _sniff_delimiter(text: str) -> str:
    sample_lines = [line for line in text.splitlines()[:20] if line.strip()][:10]
    sample = "\n".join(sample_lines)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (",", ";", "\t")}
        best = max(counts, key=counts.get)
        return best if counts[best] else ","


def _read_csv(data: bytes) -> list[tuple[int, list[str]]]:
    """All records as (record_number, cells). Quoted line breaks stay inside a cell."""
    text = _decode_csv(data)
    delimiter = _sniff_delimiter(text)
    records: list[tuple[int, list[str]]] = []
    try:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        for number, record in enumerate(reader, start=1):
            records.append((number, [to_text(cell) for cell in record]))
    except csv.Error as exc:
        raise ImportFileError(
            "bad_csv",
            "The CSV file could not be read; it may have an unclosed quote or a cell that is too large.",
            "Open it in Excel or a text editor, fix the broken row, and save it again as .csv.",
        ) from exc
    return records


def _read_xlsx(data: bytes, header_row: int, sheet_name: str | None, max_rows: int):
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency missing
        raise ImportFileError(
            "xlsx_unavailable",
            "Excel files can't be read on this server right now.",
            "Save the file as .csv and upload that instead, or ask an administrator to install openpyxl.",
        ) from exc

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - zip, XML and password errors all mean 'unreadable'
        raise ImportFileError(
            "bad_xlsx",
            "The file could not be read as an Excel workbook. It may be damaged or password protected.",
            "Open it in Excel, save a fresh copy as .xlsx (without a password), and upload that file.",
        ) from exc

    try:
        names = list(workbook.sheetnames)
        if sheet_name is not None:
            if sheet_name not in names:
                raise ImportFileError(
                    "sheet_not_found",
                    f"The sheet '{sheet_name}' was not found in this workbook.",
                    "Choose one of: " + ", ".join(names) + ".",
                )
            sheet = workbook[sheet_name]
        else:
            sheet = next((workbook[n] for n in names if hasattr(workbook[n], "iter_rows")), None)
            if sheet is None:
                raise ImportFileError(
                    "no_sheets",
                    "The workbook has no data sheets.",
                    "Add your data to a normal sheet and upload the file again.",
                )
        chosen = sheet.title

        records: list[tuple[int, list[str]]] = []
        blank_run = 0
        data_rows = 0
        for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [to_text(cell) for cell in row]
            records.append((number, cells))
            if any(c.strip() for c in cells):
                blank_run = 0
                if number > header_row:
                    data_rows += 1
                    if data_rows > max_rows:
                        break  # _build raises the too-many-rows error
            else:
                blank_run += 1
                if blank_run > MAX_CONSECUTIVE_BLANK_ROWS:
                    break
        return records, names, chosen
    finally:
        workbook.close()


# --------------------------------------------------------------------------
# Shaping
# --------------------------------------------------------------------------


def _is_blank(cells: list[str]) -> bool:
    return not any(cell.strip() for cell in cells)


def _last_filled(cells: list[str]) -> int:
    for index in range(len(cells) - 1, -1, -1):
        if cells[index].strip():
            return index
    return -1


def _build(records, header_row: int, max_rows: int, sheet_names: list[str], chosen_sheet: str | None) -> ParsedFile:
    header_cells = next((cells for number, cells in records if number == header_row), None)
    if header_cells is None or _is_blank(header_cells):
        raise ImportFileError(
            "no_header_row",
            f"Row {header_row} is empty, so no column names were found.",
            "Set the header row to the row that holds your column names, or add a header row to the file.",
        )

    data_records = [(number, cells) for number, cells in records if number > header_row]
    filled = [(number, cells) for number, cells in data_records if not _is_blank(cells)]
    if not filled:
        raise ImportFileError(
            "no_data_rows",
            "The file has column names but no data rows under them.",
            "Add your records below the header row and upload the file again.",
        )
    if len(filled) > max_rows:
        raise ImportFileError(
            "too_many_rows",
            f"The file has more than {max_rows:,} data rows.",
            f"Split it into files of {max_rows:,} rows or fewer and import them one by one.",
        )

    warnings: list[str] = []

    # Width = last column that holds a header or any data (ignores the long
    # run of empty formatted columns Excel files often have).
    width = _last_filled(header_cells) + 1
    for _, cells in filled:
        width = max(width, _last_filled(cells) + 1)

    headers: list[str] = []
    seen: dict[str, int] = {}
    for index in range(width):
        raw = header_cells[index].strip() if index < len(header_cells) else ""
        raw = " ".join(raw.split())
        if not raw:
            name = f"Column {index + 1}"
        else:
            name = raw
        key = name.lower()
        if key in seen:
            seen[key] += 1
            renamed = f"{name} ({seen[key]})"
            warnings.append(f"Two columns were both named '{name}'; the later one is shown as '{renamed}'.")
            name = renamed
            seen[name.lower()] = 1
        else:
            seen[key] = 1
        headers.append(name)

    blank_header_count = sum(1 for index in range(width) if not (header_cells[index].strip() if index < len(header_cells) else ""))
    if blank_header_count:
        warnings.append(
            f"{blank_header_count} column(s) had no name and were called 'Column N' after their position."
        )

    rows = []
    for number, cells in filled:
        padded = cells + [""] * (width - len(cells))
        rows.append(
            {
                "row_number": number,
                "values": {headers[i]: padded[i].strip() for i in range(width)},
            }
        )

    skipped_blank = len(data_records) - len(filled)
    if skipped_blank:
        warnings.append(f"{skipped_blank} completely blank row(s) were skipped.")

    return ParsedFile(
        headers=headers,
        rows=rows,
        total_rows=len(rows),
        sheet_names=sheet_names,
        warnings=warnings,
        header_row=header_row,
        sheet_name=chosen_sheet,
    )
