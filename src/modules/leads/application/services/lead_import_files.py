"""The downloadable files of a finished lead import.

Pure Python plus ``openpyxl``: no Django and no database access. Rows arrive
as iterators, so a 10,000-row file is written in small pieces:

* CSV: UTF-8 with a byte-order mark (so Excel shows accents correctly), written
  in chunks of about 64 KB;
* XLSX: openpyxl write-only mode into a temporary file that is read back in
  chunks, so the workbook is never held in memory.

Every cell that starts with ``=``, ``+``, ``-`` or ``@`` gets a leading single
quote, so a spreadsheet can never run it as a formula.
"""

from __future__ import annotations

import csv
import io
import os
import re
import tempfile
from dataclasses import dataclass
from typing import Callable, Iterable, Iterator

from .lead_import_report import (
    RESULT_LABELS,
    entry_category,
    is_rejected,
    owner_name_at,
    problem_text,
    result_details,
)

FORMAT_CSV = "csv"
FORMAT_XLSX = "xlsx"
FORMATS = (FORMAT_CSV, FORMAT_XLSX)

CONTENT_TYPES = {
    FORMAT_CSV: "text/csv; charset=utf-8",
    FORMAT_XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

PROBLEM_COLUMN = "Problem"
FIX_COLUMN = "How to fix"
RESULT_COLUMN = "Result"
DETAILS_COLUMN = "Details"
OWNER_COLUMN = "Owner"

BOM = b"\xef\xbb\xbf"
CHUNK_BYTES = 64 * 1024
MAX_CELL_LENGTH = 32_000  # Excel allows 32,767 characters in a cell
_FORMULA_STARTS = ("=", "+", "-", "@")
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------

def safe_cell(value) -> str:
    """Text for one written cell: never a formula, never a control character."""
    if value is None:
        return ""
    text = value if isinstance(value, str) else str(value)
    text = _ILLEGAL_XML.sub("", text)
    if len(text) > MAX_CELL_LENGTH:
        text = text[:MAX_CELL_LENGTH]
    if text.startswith(_FORMULA_STARTS):
        text = "'" + text
    return text


def unique_names(headers: list[str], wanted: list[str]) -> list[str]:
    """The extra column names, changed to ``Name (2)`` when a column of the
    uploaded file already uses that name (for example a rejected file that was
    uploaded again with its Problem column left in)."""
    taken = {str(h).strip().lower() for h in headers}
    names = []
    for name in wanted:
        candidate, number = name, 1
        while candidate.lower() in taken:
            number += 1
            candidate = f"{name} ({number})"
        taken.add(candidate.lower())
        names.append(candidate)
    return names


def rejected_columns(headers: list[str]) -> list[str]:
    return list(headers) + unique_names(headers, [PROBLEM_COLUMN, FIX_COLUMN])


def results_columns(headers: list[str]) -> list[str]:
    return list(headers) + unique_names(headers, [RESULT_COLUMN, DETAILS_COLUMN, OWNER_COLUMN])


# ---------------------------------------------------------------------------
# File name
# ---------------------------------------------------------------------------

def base_name(file_name: str) -> str:
    """The uploaded file's name without its extension and without path parts."""
    name = os.path.basename(str(file_name or "").replace("\\", "/")).strip()
    stem = os.path.splitext(name)[0].strip().strip(".")
    stem = re.sub(r"[\x00-\x1f\"]", "", stem)
    return stem[:120] or "import"


def download_name(file_name: str, suffix: str, file_format: str) -> str:
    return f"{base_name(file_name)}-{suffix}.{file_format}"


def default_format(file_name: str) -> str:
    """The uploaded file's own format when it is one of ours, else CSV."""
    extension = os.path.splitext(str(file_name or ""))[1].lower().lstrip(".")
    return extension if extension in FORMATS else FORMAT_CSV


# ---------------------------------------------------------------------------
# Rows
# ---------------------------------------------------------------------------

def _values_row(headers: list[str], values: dict | None) -> list[str]:
    values = values or {}
    return [safe_cell(values.get(header)) for header in headers]


class _RowCursor:
    """Finds stored rows by row number when the pages arrive in file order and
    the wanted numbers are asked for in ascending order."""

    def __init__(self, pages: Iterable[list[dict]]):
        self._rows = (row for page in pages for row in page)
        self._current = None

    def values_of(self, row_number: int) -> dict | None:
        while True:
            if self._current is None:
                self._current = next(self._rows, None)
                if self._current is None:
                    return None
            number = self._current.get("row_number", 0)
            if number < row_number:
                self._current = None
                continue
            if number == row_number:
                row, self._current = self._current, None
                return row.get("values", row.get("original")) or {}
            return None  # the row is not in the file any more


def rejected_rows(
    headers: list[str],
    entries: list[dict],
    page_source: Callable[[], Iterable[list[dict]]],
    include_excluded: bool = False,
) -> Iterator[list[str]]:
    """The cells of each rejected row, in file order: the original values,
    then Problem and How to fix.

    Values come from the entry when it kept them; when the file had too many
    problem rows to keep them, they are read back from the stored rows (one
    pass over the pages, in order).
    """
    chosen = sorted(
        (e for e in entries if is_rejected(e, include_excluded)),
        key=lambda e: e.get("row_number", 0),
    )
    cursor = None
    for entry in chosen:
        values = entry.get("original_values")
        if values is None:
            if cursor is None:
                cursor = _RowCursor(page_source())
            values = cursor.values_of(entry.get("row_number", 0))
        problem, fix = problem_text(entry)
        yield _values_row(headers, values) + [safe_cell(problem), safe_cell(fix)]


def count_rejected(entries: list[dict], include_excluded: bool = False) -> int:
    return sum(1 for entry in entries if is_rejected(entry, include_excluded))


def results_rows(
    headers: list[str],
    entries: list[dict],
    pages: Iterable[list[dict]],
    options: dict,
) -> Iterator[list[str]]:
    """One row per uploaded row: the original values, then Result, Details and
    Owner. A row without a stored entry was a plainly added lead."""
    by_number = {e.get("row_number"): e for e in entries}
    position = 0
    for page in pages:
        for row in page:
            number = row.get("row_number")
            entry = by_number.get(number)
            category = entry_category(entry) if entry else "added"
            owner = (entry or {}).get("owner_name") or owner_name_at(options, position)
            position += 1
            yield _values_row(headers, row.get("values", row.get("original"))) + [
                safe_cell(RESULT_LABELS[category]),
                safe_cell(result_details(entry)),
                safe_cell(owner),
            ]


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

@dataclass
class FileStream:
    """A file that is read out in chunks. ``close`` removes any temporary file."""

    chunks: Iterator[bytes]
    content_type: str
    close: Callable[[], None] = lambda: None

    def __iter__(self):
        try:
            yield from self.chunks
        finally:
            self.close()


def csv_stream(columns: list[str], rows: Iterable[list[str]]) -> FileStream:
    def generate():
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, lineterminator="\r\n")
        yield BOM
        writer.writerow([safe_cell(column) for column in columns])
        for row in rows:
            writer.writerow(row)
            if buffer.tell() >= CHUNK_BYTES:
                yield buffer.getvalue().encode("utf-8")
                buffer.seek(0)
                buffer.truncate(0)
        if buffer.tell():
            yield buffer.getvalue().encode("utf-8")

    return FileStream(generate(), CONTENT_TYPES[FORMAT_CSV])


def xlsx_stream(
    columns: list[str],
    rows: Iterable[list[str]],
    sheet_title: str,
    wide_columns: tuple[str, ...] = (),
) -> FileStream:
    """Builds the workbook first (write-only, so it stays small), then reads it
    back in chunks."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    handle = tempfile.NamedTemporaryFile(prefix="lead-import-", suffix=".xlsx", delete=False)
    path = handle.name
    handle.close()

    def cleanup():
        try:
            os.remove(path)
        except OSError:
            pass

    try:
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet(title=sheet_title[:31])
        sheet.freeze_panes = "A2"
        for index, column in enumerate(columns, start=1):
            width = 55 if column in wide_columns else 20
            sheet.column_dimensions[get_column_letter(index)].width = width

        header = []
        for column in columns:
            cell = WriteOnlyCell(sheet, value=safe_cell(column))
            cell.font = Font(bold=True)
            header.append(cell)
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
    except Exception:
        cleanup()
        raise

    def generate():
        with open(path, "rb") as source:
            while True:
                chunk = source.read(CHUNK_BYTES)
                if not chunk:
                    break
                yield chunk

    return FileStream(generate(), CONTENT_TYPES[FORMAT_XLSX], close=cleanup)


def build_stream(
    file_format: str,
    columns: list[str],
    rows: Iterable[list[str]],
    sheet_title: str,
    wide_columns: tuple[str, ...] = (),
) -> FileStream:
    if file_format == FORMAT_XLSX:
        return xlsx_stream(columns, rows, sheet_title, wide_columns)
    return csv_stream(columns, rows)


__all__ = [
    "FORMATS",
    "FileStream",
    "build_stream",
    "count_rejected",
    "default_format",
    "download_name",
    "rejected_columns",
    "rejected_rows",
    "results_columns",
    "results_rows",
    "safe_cell",
]
